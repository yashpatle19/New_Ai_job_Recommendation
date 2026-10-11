import asyncio
import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException
from starlette.requests import Request
from docx import Document

from backend.main import post_recommend
from backend.recommend import load_jobs, recommend_jobs
from backend.resume_parser import extract_resume_text


class RecommendationTests(unittest.TestCase):
    def test_relevant_job_ranks_first_and_explains_skills(self):
        rows = recommend_jobs("Python machine learning pandas scikit-learn", top_k=2)
        self.assertEqual(rows.iloc[0].title, "Machine Learning Engineer")
        self.assertIn("python", rows.iloc[0].matched_skills)
        self.assertTrue(rows.iloc[0].missing_skills)
        self.assertIn("0.7", rows.iloc[0].explanation)

    def test_empty_resume_rejected(self):
        with self.assertRaises(ValueError):
            recommend_jobs("  ")

    def test_missing_optional_columns_filled(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "jobs.csv"
            path.write_text("title\nEngineer\n", encoding="utf-8")
            jobs = load_jobs(path)
            self.assertEqual(jobs.iloc[0].skills, "")
            self.assertEqual(jobs.iloc[0].company, "")

    def test_missing_title_column_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "jobs.csv"
            path.write_text("role\nEngineer\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "title"):
                load_jobs(path)

    def test_invalid_file_format_rejected_before_read(self):
        with self.assertRaisesRegex(ValueError, "Unsupported"):
            extract_resume_text("resume.txt", b"resume")

    def test_docx_resume_extraction(self):
        document = Document()
        document.add_paragraph("Python machine learning pandas")
        buffer = BytesIO()
        document.save(buffer)
        self.assertIn("Python", extract_resume_text("resume.docx", buffer.getvalue()))

    def test_recommend_api_response_shape(self):
        body = json.dumps({"resume_text": "Python SQL pandas", "top_k": 2}).encode()
        request = Request({"type": "http", "method": "POST", "path": "/recommend", "headers": [(b"content-type", b"application/json")], "query_string": b"", "server": ("test", 80), "client": ("test", 1), "scheme": "http"}, receive=lambda: None)
        async def receive():
            return {"type": "http.request", "body": body, "more_body": False}
        request._receive = receive
        response = asyncio.run(post_recommend(request))
        self.assertEqual(response["count"], 2)
        self.assertIn("scoring_formula", response)
        self.assertIn("match_score", response["recommendations"][0])

    def test_recommend_api_rejects_empty_resume(self):
        body = b'{"resume_text":"  ","top_k":1}'
        request = Request({"type": "http", "method": "POST", "path": "/recommend", "headers": [(b"content-type", b"application/json")], "query_string": b"", "server": ("test", 80), "client": ("test", 1), "scheme": "http"})
        async def receive():
            return {"type": "http.request", "body": body, "more_body": False}
        request._receive = receive
        with self.assertRaises(HTTPException) as error:
            asyncio.run(post_recommend(request))
        self.assertEqual(error.exception.status_code, 422)

    def test_multipart_docx_upload(self):
        document = Document()
        document.add_paragraph("Python SQL pandas")
        buffer = BytesIO()
        document.save(buffer)
        boundary = "test-boundary"
        body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"top_k\"\r\n\r\n2\r\n"
                f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"resume.docx\"\r\n"
                "Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n").encode()
        body += buffer.getvalue() + f"\r\n--{boundary}--\r\n".encode()
        async def receive():
            return {"type": "http.request", "body": body, "more_body": False}
        request = Request({"type": "http", "method": "POST", "path": "/recommend", "headers": [(b"content-type", f"multipart/form-data; boundary={boundary}".encode()), (b"content-length", str(len(body)).encode())], "query_string": b"", "server": ("test", 80), "client": ("test", 1), "scheme": "http"}, receive=receive)
        response = asyncio.run(post_recommend(request))
        self.assertEqual(response["count"], 2)


if __name__ == "__main__":
    unittest.main()
